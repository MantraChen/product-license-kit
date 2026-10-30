"""Export the guide with python-docx; --check verifies source and body parity.

Run in a documentation environment with python-docx. Render and visually review
the output before committing. Unsupported Markdown blocks fail explicitly.
"""
import argparse
import hashlib
import re
from datetime import datetime, timezone
from pathlib import Path
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'docs/SYSTEM-GUIDE.md'
OUTPUT = SOURCE.with_suffix('.docx')


def plain(text):
    return re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', text).replace('`', '').replace('**', '')


def parse(text):
    lines = iter(text.splitlines())
    table = []
    for raw in [*lines, '']:
        line = raw.strip()
        if line.startswith('|'):
            table.append([cell.strip() for cell in line.strip('|').split('|')])
            continue
        if table:
            if len(table) < 2 or not all(re.fullmatch(r':?-+:?', cell) for cell in table[1]):
                raise ValueError('Invalid table header')
            rows = [table[0], *table[2:]]
            if any(len(row) != len(rows[0]) for row in rows):
                raise ValueError('Invalid table width')
            yield 'table', rows
            table = []
        if not line:
            continue
        if line == '<!-- pagebreak -->':
            yield 'break', None
        elif re.match(r'^#{1,3} ', line):
            head, content = line.split(' ', 1)
            yield 'heading', (len(head), content)
        elif line.startswith('- '):
            yield 'paragraph', '• ' + line[2:]
        elif line.startswith(('```', '<!--', '>', '![', '#')):
            raise ValueError('Unsupported block: ' + line[:40])
        else:
            yield 'paragraph', line


def set_font(style, size):
    style.font.name = 'Microsoft YaHei'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for kind in ('ascii', 'hAnsi', 'eastAsia'):
        fonts.set(qn('w:' + kind), 'Microsoft YaHei')


def write_text(paragraph, text):
    for part in re.split(r'(\[[^\]]+\]\([^)]+\)|`[^`]+`|\*\*[^*]+\*\*)', text):
        if not part:
            continue
        if part.startswith('[') and '](' in part:
            label, target = part[1:-1].split('](', 1)
            if not target.startswith(('https://', 'http://')):
                relative = (SOURCE.parent / target).resolve().relative_to(ROOT)
                target = 'https://github.com/MantraChen/' + ROOT.name + '/blob/main/' + relative.as_posix()
            link = OxmlElement('w:hyperlink')
            link.set(qn('r:id'), paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True))
            run = OxmlElement('w:r')
            content = OxmlElement('w:t')
            content.text = label
            run.append(content)
            link.append(run)
            paragraph._p.append(link)
            continue
        run = paragraph.add_run(plain(part))
        if part.startswith('**'):
            run.bold = True
        if part.startswith('`'):
            run.font.name = 'Consolas'
            run.font.size = Pt(9)


def table_block(doc, rows):
    widths = {3: [3.0, 6.0, 8.4], 4: [3.6, 4.6, 4.6, 4.6]}[len(rows[0])]
    if rows[0] == ['角色', '已有能力', '限制']:
        widths = [3.0, 7.4, 7.0]
    table = doc.add_table(rows=0, cols=len(widths))
    table.autofit = False
    for column, width in zip(table.columns, widths):
        column.width = Cm(width)
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge = OxmlElement('w:' + side)
        for key, value in (('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')):
            edge.set(qn('w:' + key), value)
        borders.append(edge)
    table._tbl.tblPr.append(borders)
    for index, values in enumerate(rows):
        row = table.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if index == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for cell, text, width in zip(row.cells, values, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side, value in (('top', '85'), ('bottom', '85'), ('left', '100'), ('right', '100')):
                edge = OxmlElement('w:' + side)
                edge.set(qn('w:w'), value)
                edge.set(qn('w:type'), 'dxa')
                margins.append(edge)
            pr.append(margins)
            if index == 0:
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'), 'E9EEF2')
                pr.append(shade)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            write_text(p, text)
            for run in p.runs:
                run.font.size = Pt(9)
                run.bold = index == 0
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(0)
    spacer.paragraph_format.line_spacing = Pt(4)


def verify(blocks, digest):
    doc = Document(OUTPUT)
    if doc.core_properties.identifier != digest:
        raise ValueError('DOCX source digest differs; regenerate')
    expected = []
    for kind, value in blocks:
        if kind == 'table':
            expected.extend(plain(cell) for row in value for cell in row)
        elif kind == 'heading':
            expected.append(plain(value[1]))
        elif kind == 'paragraph':
            expected.append(plain(value))
    actual = ''.join(node.text or '' for node in doc.element.body.iter(qn('w:t')))
    if actual != ''.join(expected):
        raise ValueError('DOCX body differs from Markdown')
    print(f'Synced: {OUTPUT.name}; source SHA-256 {digest}')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    # Normalize line endings so a Git checkout on another platform stays valid.
    text = SOURCE.read_text(encoding='utf-8-sig')
    digest = hashlib.sha256(text.encode('utf-8')).hexdigest()
    blocks = list(parse(text))
    if args.check:
        verify(blocks, digest)
        return
    doc = Document()
    # The bundled template may carry a blue rule on Title; remove it explicitly.
    for border in doc.styles.element.xpath('.//w:pBdr'):
        border.getparent().remove(border)
    # Remove the template's empty bibliography custom XML from the deliverable.
    for rel_id, rel in list(doc.part.rels.items()):
        if rel.reltype == RELATIONSHIP_TYPE.CUSTOM_XML:
            doc.part.drop_rel(rel_id)
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(1.8), Cm(1.7)
    section.left_margin = section.right_margin = Cm(1.8)
    section.footer_distance = Cm(0.7)
    normal = doc.styles['Normal']
    set_font(normal, 10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.2
    normal.paragraph_format.widow_control = True
    for name, size in (('Title', 22), ('Heading 1', 16), ('Heading 2', 12)):
        style = doc.styles[name]
        set_font(style, size)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(10 if name != 'Title' else 0)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    footer.add_run(' 页')
    for kind, value in blocks:
        if kind == 'break':
            doc.add_page_break()
        elif kind == 'table':
            table_block(doc, value)
        else:
            if kind == 'heading':
                level, value = value
                p = doc.add_paragraph(style={1: 'Title', 2: 'Heading 1', 3: 'Heading 2'}[level])
            else:
                p = doc.add_paragraph()
            write_text(p, value)
    props = doc.core_properties
    props.title = blocks[0][1][1]
    props.author = props.last_modified_by = 'MantraChen'
    props.identifier = digest
    props.comments = 'Generated from docs/SYSTEM-GUIDE.md; edit the source and regenerate.'
    props.created = props.modified = datetime(2026, 10, 31, tzinfo=timezone.utc)
    doc.save(OUTPUT)
    verify(blocks, digest)


if __name__ == '__main__':
    main()
